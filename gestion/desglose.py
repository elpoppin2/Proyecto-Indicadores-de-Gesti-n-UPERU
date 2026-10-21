"""Detalle fila por fila de los indicadores de cada categoría (botón "Ver
detalle fila por fila" al pie de cada sección). Pedido por el usuario: ir de
lo general a lo específico y ver de qué filas de SAP sale cada número (líneas
de pedido ME2N, movimientos MB51, inventario MB5B, tickets ZMM056...).

Cada detalle se define aquí (columnas, filtros por clic, resumen) y sus filas
las arma views.py con los mismos segmentadores que la sección (período,
corte, familia, centro...). Este módulo filtra, busca, ordena, pagina y
exporta a Excel. El resumen se calcula sobre las filas que quedan después de
los filtros, así que sin filtros cuadra con el indicador de la sección.
"""
import collections
import datetime
import io
import math
import threading
from urllib.parse import urlencode

import pandas as pd
from django.http import Http404, HttpResponse
from django.template.loader import render_to_string

from dashboard import queries as q

# Filas por página (el resto se pide con "Cargar más" o al bajar el scroll)
# y tope de la descarga a Excel.
FILAS_POR_PAGINA = 200
MAX_FILAS_EXCEL = 150_000

# Parámetros propios del detalle (no son segmentadores de la página).
PARAMETROS_UI = {"q", "orden", "dir", "offset", "_parte", "formato", "_zona", "_seccion"}
# Filtros por clic que cambian la consulta base (el período), no solo filtran filas.
FILTROS_DE_PERIODO = {"f_mes", "f_anio"}

_TIPOS_NUMERICOS = {"soles", "numero", "cantidad", "entero", "dias", "pct"}
_TIPOS_TEXTO = {"texto", "codigo", "centro"}


def _col(clave, titulo, tipo="texto", tono=None):
    return {"clave": clave, "titulo": titulo, "tipo": tipo, "tono": tono, "num": tipo in _TIPOS_NUMERICOS}


def _tono_signo(valor, inverso=False):
    """Verde si sube / rojo si baja (o al revés con `inverso`: un precio que sube es malo)."""
    if valor is None or (isinstance(valor, float) and math.isnan(valor)) or valor == 0:
        return ""
    malo = valor > 0 if inverso else valor < 0
    return "dato-negativo" if malo else "dato-positivo"


def _tono_estado(valor):
    if valor in ("Cumple", "OK", "Completa", "Válida", "Sí"):
        return "dato-positivo"
    if valor in ("No cumple", "En riesgo", "Incompleta", "No"):
        return "dato-negativo"
    return "texto-tenue" if valor else ""


# --- Resúmenes: (etiqueta, texto) sobre las filas filtradas ------------------

def _suma(df, col):
    return float(df[col].sum()) if col in df and not df.empty else 0.0


def _pct(parte, total):
    return parte / total * 100 if total else None


def _texto_pct(valor, decimales=1):
    return "—" if valor is None else f"{valor:,.{decimales}f}%"


def _soles(valor):
    return f"S/ {valor:z,.0f}"


def _pmp(df):
    valor = _suma(df, "valor")
    return float((df["dias"] * df["valor"]).sum() / valor) if valor else None


def _resumen_compras(df):
    valor = _suma(df, "valor")
    contrato = float(df.loc[df["bajo_contrato"], "valor"].sum()) if not df.empty else 0.0
    pmp = _pmp(df)
    return [("Líneas", f"{len(df):,}"), ("Pedidos", f"{df['pedido'].nunique():,}"),
            ("Valor de compras", _soles(valor)), ("Bajo contrato", _texto_pct(_pct(contrato, valor))),
            ("PMP", "—" if pmp is None else f"{pmp:z,.0f} días")]


def _resumen_cuentas(df):
    valor = _suma(df, "valor")
    pmp = _pmp(df)
    a_60 = float(df.loc[df["dias"] >= 60, "valor"].sum()) if not df.empty else 0.0
    return [("Líneas", f"{len(df):,}"), ("Valor de compras", _soles(valor)),
            ("Σ días × valor", _soles(_suma(df, "peso"))),
            ("PMP = Σ días × valor ÷ Σ valor", "—" if pmp is None else f"{pmp:z,.0f} días"),
            ("A 60 días o más", _texto_pct(_pct(a_60, valor)))]


def _resumen_fillrate(df):
    valor = _suma(df, "valor")
    completas = float(df.loc[df["completa"] == "Completa", "valor"].sum()) if not df.empty else 0.0
    return [("Líneas evaluadas", f"{len(df):,}"), ("Valor", _soles(valor)),
            ("Valor entregado completo", _soles(completas)), ("Fill Rate", _texto_pct(_pct(completas, valor)))]


def _resumen_leadtime(df):
    validas = df[df["estado"] == "Válida"]
    if validas.empty:
        return [("Líneas", f"{len(df):,}")]
    return [("Líneas válidas", f"{len(validas):,}"), ("Lead Time promedio", f"{validas['dias'].mean():z,.1f} días"),
            ("Mediana", f"{validas['dias'].median():z,.1f} días"),
            ("Excluidas", f"{len(df) - len(validas):,}")]


def _resumen_otif(df):
    evaluadas = df[df["estado"] != "En plazo"]
    valor = _suma(evaluadas, "valor")

    def pct(columna):
        return _texto_pct(_pct(float(evaluadas.loc[evaluadas[columna], "valor"].sum()), valor)) if valor else "—"
    return [("Líneas evaluadas", f"{len(evaluadas):,}"), ("Valor evaluado", _soles(valor)),
            ("OTIF", pct("otif")), ("A tiempo", pct("a_tiempo")), ("Completo", pct("completo")),
            ("Todavía en plazo", f"{len(df) - len(evaluadas):,}")]


def _resumen_unico(df):
    valor = _suma(df, "valor")
    unicos = df[df["un_solo"]]
    return [("Materiales recurrentes", f"{len(df):,}"), ("Con un solo proveedor", f"{len(unicos):,}"),
            ("Valor recurrente", _soles(valor)), ("Con proveedor único", _soles(_suma(unicos, "valor"))),
            ("% del valor", _texto_pct(_pct(_suma(unicos, "valor"), valor), 0))]


def _resumen_inventario(df):
    return [("Materiales", f"{len(df):,}"), ("Valor", _soles(_suma(df, "valor_total"))),
            ("Clase A", f"{int((df['clase'] == 'A').sum()):,}"), ("Clase B", f"{int((df['clase'] == 'B').sum()):,}"),
            ("Clase C", f"{int((df['clase'] == 'C').sum()):,}")]


def _resumen_sinmov(df):
    nunca = int(df["ultimo_consumo"].isna().sum()) if not df.empty else 0
    return [("Materiales sin movimiento", f"{len(df):,}"), ("Valor inmovilizado", _soles(_suma(df, "valor_total"))),
            ("Sin ningún consumo en la data cargada", f"{nunca:,}")]


def _resumen_consumo(df):
    return [("Movimientos", f"{len(df):,}"), ("Materiales", f"{df['material_codigo'].nunique():,}"),
            ("Consumo valorizado", _soles(_suma(df, "importe")))]


def _resumen_criticos(df):
    con = df[df["estado"] != "Sin consumo"]
    return [("Críticos con cobertura", f"{len(con):,}"),
            ("Cobertura promedio", "—" if con.empty else f"{con['cobertura_meses'].mean():z,.2f} meses"),
            ("En riesgo (< 1 mes)", f"{int((df['estado'] == 'En riesgo').sum()):,}"),
            ("Sin consumo (no entran al promedio)", f"{len(df) - len(con):,}")]


def _resumen_variacion(df):
    return [("Materiales", f"{len(df):,}"), ("Valor inicial", _soles(_suma(df, "valor_mes_inicial"))),
            ("Valor final", _soles(_suma(df, "valor_mes_final"))),
            ("Variación", _soles(_suma(df, "variacion_valor")))]


def _resumen_precios(df):
    validas = df[~df["atipica"]]
    base = float((validas["precio_base"] * validas["toneladas"]).sum())
    importe = _suma(validas, "importe")
    return [("Comparaciones", f"{len(df):,}"),
            ("Materiales", f"{df[['material_codigo', 'centro']].drop_duplicates().shape[0]:,}"),
            ("Pagado", _soles(importe)), ("Al precio base", _soles(base)),
            ("Variación ponderada", _texto_pct((importe / base - 1) * 100 if base else None)),
            ("Bolsa", _soles(importe - base)), ("Atípicas (no suman)", f"{len(df) - len(validas):,}")]


def _resumen_entradas(df):
    return [("Movimientos", f"{len(df):,}"), ("Materiales", f"{df['material_codigo'].nunique():,}"),
            ("Importe", _soles(_suma(df, "importe")))]


def _resumen_abastecimiento(df):
    plan, real = _suma(df, "cantidad_plan"), _suma(df, "cantidad_real")
    return [("Materiales", f"{len(df):,}"), ("Plan", f"{plan:z,.0f} T"), ("Real", f"{real:z,.0f} T"),
            ("Cumplimiento", _texto_pct(_pct(real, plan)))]


def _resumen_ingresos(df):
    return [("Tickets", f"{len(df):,}"), ("Toneladas", f"{_suma(df, 'toneladas'):z,.0f} T"),
            ("Proveedores u orígenes", f"{df['proveedor'].nunique():,}")]


# --- Definiciones -----------------------------------------------------------
# columnas: (clave, título, tipo); filtros: parámetro GET -> columna (clic en
# una fila, barra o punto de la sección; también son las "chips" quitables).

_COL_MATERIAL = _col("material_codigo", "Material", "codigo")
_COL_DESCRIPCION = _col("texto_breve", "Descripción")

DESGLOSES = {
    "compras": {
        "titulo": "Líneas de pedido",
        "nota": "Una fila por posición de pedido (ME2N) con Fecha doc. en el período: Valor (S/) = Tp. cambio × Valor neto. "
                "Bajo contrato: pedido 47…/40… o con contrato (46…). Días de pago según la condición de pago (CPag).",
        "columnas": [
            _col("fecha_documento", "Fecha doc.", "fecha"), _col("pedido", "Pedido", "codigo"),
            _col("posicion", "Pos.", "codigo"), _col("clase_documento", "Cl. doc.", "codigo"),
            _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedor", "Proveedor"), _col("centro", "Centro", "centro"),
            _col("cantidad_pedido", "Cantidad", "cantidad"), _col("unidad", "UM"),
            _col("precio_neto", "Precio neto", "numero"), _col("moneda", "Mon."),
            _col("valor_neto", "Valor neto", "numero"), _col("tipo_cambio", "Tp. cambio", "numero"),
            _col("valor", "Valor (S/)", "soles"), _col("contrato_texto", "Bajo contrato"),
            _col("condicion_pago", "CPag", "codigo"), _col("dias", "Días de pago", "dias"),
        ],
        "orden": ("valor", "desc"),
        "filtros": {"f_proveedor": "proveedor", "f_material": "material_codigo", "f_contrato": "contrato_si"},
        "resumen": _resumen_compras,
    },
    "cuentas": {
        "titulo": "Líneas de pedido con su plazo de pago",
        "nota": "Una fila por posición de pedido (ME2N) con Fecha doc. en el período. Días de pago: los de su condición de "
                "pago (CPag). PMP = Σ (días × valor) ÷ Σ valor.",
        "columnas": [
            _col("fecha_documento", "Fecha doc.", "fecha"), _col("pedido", "Pedido", "codigo"),
            _col("posicion", "Pos.", "codigo"), _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedor", "Proveedor"),
            _col("moneda", "Mon."), _col("valor_neto", "Valor neto", "numero"),
            _col("tipo_cambio", "Tp. cambio", "numero"), _col("valor", "Valor (S/)", "soles"),
            _col("condicion_pago", "CPag", "codigo"), _col("dias", "Días de pago", "dias"),
            _col("tramo", "Plazo"), _col("peso", "Días × valor", "soles"),
        ],
        "orden": ("valor", "desc"),
        "filtros": {"f_proveedor": "proveedor", "f_condicion": "condicion_pago", "f_tramo": "tramo"},
        "resumen": _resumen_cuentas,
    },
    "fillrate": {
        "titulo": "Líneas de material e importación (ZNAC/ZIMP)",
        "nota": "Una fila por posición de pedido ZNAC/ZIMP con Fecha doc. en el período y cantidad pedida mayor a 0. "
                "Completa: lo recibido llega al 95% de lo pedido o más. Fill Rate = valor de las completas ÷ valor total.",
        "columnas": [
            _col("fecha_documento", "Fecha doc.", "fecha"), _col("pedido", "Pedido", "codigo"),
            _col("posicion", "Pos.", "codigo"), _col("clase_documento", "Cl. doc.", "codigo"),
            _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedor", "Proveedor"),
            _col("cantidad_pedido", "Pedido (cant.)", "cantidad"), _col("cantidad_entrada", "Recibido", "cantidad"),
            _col("unidad", "UM"), _col("pct_entregado", "% recibido", "pct"),
            _col("completa", "Estado", "texto", _tono_estado), _col("valor", "Valor (S/)", "soles"),
        ],
        "orden": ("valor", "desc"),
        "filtros": {"f_proveedor": "proveedor", "f_estado": "completa"},
        "resumen": _resumen_fillrate,
    },
    "otif": {
        "titulo": "Líneas evaluadas en el OTIF",
        "nota": "Una fila por posición de pedido ZNAC/ZIMP con fecha de entrega prometida en el período. Fecha límite = "
                "prometida + días de tolerancia. Recibido a tiempo = entradas (MB51 101 − 102) hasta la fecha límite. "
                "Las líneas todavía en plazo no entran al cálculo.",
        "columnas": [
            _col("pedido", "Pedido", "codigo"), _col("posicion", "Pos.", "codigo"), _COL_MATERIAL, _COL_DESCRIPCION,
            _col("proveedor", "Proveedor"), _col("fecha_entrega", "Fecha prometida", "fecha"),
            _col("limite", "Fecha límite", "fecha"), _col("primera_entrada", "Primera entrada", "fecha"),
            _col("ultima_entrada", "Última entrada", "fecha"), _col("cantidad_pedido", "Pedido (cant.)", "cantidad"),
            _col("recibido_a_tiempo", "Recibido a tiempo", "cantidad"), _col("recibido_total", "Recibido total", "cantidad"),
            _col("a_tiempo_texto", "A tiempo", "texto", _tono_estado), _col("completo_texto", "Completo", "texto", _tono_estado),
            _col("estado", "OTIF", "texto", _tono_estado), _col("motivo", "Motivo"), _col("valor", "Valor (S/)", "soles"),
        ],
        "orden": ("valor", "desc"),
        "filtros": {"f_proveedor": "proveedor", "f_estado": "estado"},
        "resumen": _resumen_otif,
    },
    "leadtime": {
        "titulo": "Líneas con entrada de mercancía",
        "nota": "Una fila por posición de pedido ZNAC/ZIMP con Fecha doc. en el período que ya tuvo entrada de mercancía "
                "(MB51 101). Lead Time = primera entrada − fecha del pedido. Se excluyen las entradas anteriores al pedido.",
        "columnas": [
            _col("pedido", "Pedido", "codigo"), _col("posicion", "Pos.", "codigo"), _col("clase_documento", "Cl. doc.", "codigo"),
            _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedor", "Proveedor"), _col("fecha_pedido", "Fecha pedido", "fecha"),
            _col("fecha_entrega", "Fecha prometida", "fecha"), _col("fecha_recepcion", "Primera entrada", "fecha"),
            _col("dias", "Lead Time (días)", "dias"), _col("estado", "Estado", "texto", _tono_estado),
            _col("valor", "Valor (S/)", "soles"),
        ],
        "orden": ("dias", "desc"),
        "filtros": {"f_proveedor": "proveedor", "f_estado": "estado"},
        "resumen": _resumen_leadtime,
    },
    "unico": {
        "titulo": "Materiales recurrentes y sus proveedores",
        "nota": "Una fila por material comprado de forma recurrente en el período (sale de las líneas de pedido ME2N): "
                "cuántos proveedores tuvo, cuáles, en cuántos meses se compró y por cuánto.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedores", "N.º proveedores", "entero"),
            _col("nombres", "Proveedores"), _col("meses", "Meses con compra", "entero"),
            _col("pedidos", "Pedidos", "entero"), _col("valor", "Valor (S/)", "soles"),
            _col("un_solo_texto", "Un solo proveedor", "texto", _tono_estado),
        ],
        "orden": ("valor", "desc"),
        "filtros": {"f_unico": "un_solo_texto", "f_material": "material_codigo"},
        "resumen": _resumen_unico,
    },
    "inventario": {
        "titulo": "Materiales con stock al corte",
        "nota": "Una fila por material del inventario al corte (MB5B: stock y valor de cierre). Clase ABC: materiales "
                "ordenados por valor, A hasta el 80% acumulado, B hasta el 95%, C el resto.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("familia", "Familia"), _col("centro", "Centro", "centro"),
            _col("unidad_medida_base", "UMB"), _col("stock_cierre", "Stock", "cantidad"),
            _col("valor_total", "Valor (S/)", "soles"), _col("pct", "% del valor", "pct"),
            _col("pct_acumulado", "% acumulado", "pct"), _col("clase", "Clase ABC"),
        ],
        "orden": ("valor_total", "desc"),
        "filtros": {"f_familia": "familia", "f_centro": "centro", "f_clase": "clase", "f_material": "material_codigo"},
        "resumen": _resumen_inventario,
    },
    "sinmov": {
        "titulo": "Materiales sin movimiento",
        "nota": "Materiales con valor en el corte (MB5B) y ningún consumo (MB51) en los 3 años que terminan en el corte. "
                "Último consumo y última entrada: en toda la historia cargada de MB51.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("familia", "Familia"), _col("centro", "Centro", "centro"),
            _col("unidad_medida_base", "UMB"), _col("stock_cierre", "Stock", "cantidad"),
            _col("valor_total", "Valor (S/)", "soles"), _col("ultimo_consumo", "Último consumo", "fecha"),
            _col("ultima_entrada", "Última entrada", "fecha"),
        ],
        "orden": ("valor_total", "desc"),
        "filtros": {"f_familia": "familia", "f_centro": "centro"},
        "resumen": _resumen_sinmov,
    },
    "consumo": {
        "titulo": "Movimientos de consumo (MB51)",
        "nota": "Una fila por movimiento de consumo de los 6 meses que terminan en el corte: clases Z01, Z04, 353 y "
                "200 a 299 (las anulaciones restan). Consumo e importe en positivo = salida de almacén. Es el "
                "consumo valorizado de la Rotación.",
        "columnas": [
            _col("fecha", "Fecha", "fecha"), _col("documento", "Doc. material", "codigo"),
            _col("posicion", "Pos.", "codigo"), _col("clase_movimiento", "Cl. mov.", "codigo"),
            _col("texto_clase", "Movimiento"), _COL_MATERIAL, _COL_DESCRIPCION,
            _col("centro_movimiento", "Centro", "centro"), _col("almacen", "Alm.", "codigo"),
            _col("consumo", "Consumo", "cantidad"), _col("unidad", "UM"), _col("importe", "Importe (S/)", "soles"),
            _col("orden", "Orden", "codigo"), _col("centro_coste", "Centro de coste", "codigo"),
            _col("lote", "Lote", "codigo"), _col("usuario", "Usuario"),
        ],
        "orden": ("importe", "desc"),
        "filtros": {"f_material": "material_codigo", "f_clase_mov": "clase_movimiento"},
        "resumen": _resumen_consumo,
    },
    "criticos": {
        "titulo": "Materiales críticos",
        "nota": "Materiales marcados como críticos en el maestro (ZMM060). Cobertura = stock del corte ÷ consumo promedio "
                "de 6 meses, en cantidad. Los que no tuvieron consumo no tienen cobertura y no entran al promedio.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("familia", "Familia"), _col("centro", "Centro", "centro"),
            _col("unidad_medida_base", "UMB"), _col("stock_cierre", "Stock", "cantidad"),
            _col("consumo_promedio", "Consumo prom. (6m)", "cantidad"), _col("cobertura_meses", "Cobertura (meses)", "numero"),
            _col("estado", "Estado", "texto", _tono_estado),
        ],
        "orden": ("cobertura_meses", "asc"),
        "filtros": {"f_estado": "estado"},
        "resumen": _resumen_criticos,
    },
    "variacion": {
        "titulo": "Variación por material",
        "nota": "Todos los materiales de los dos cortes (MB5B), no solo los 20 que más subieron.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("familia", "Familia"), _col("centro", "Centro", "centro"),
            _col("unidad_medida_base", "UMB"), _col("stock_mes_inicial", "Stock inicial", "cantidad"),
            _col("stock_mes_final", "Stock final", "cantidad"), _col("variacion_stock", "Variación stock", "cantidad"),
            _col("valor_mes_inicial", "Valor inicial (S/)", "soles"), _col("valor_mes_final", "Valor final (S/)", "soles"),
            _col("variacion_valor", "Variación (S/)", "soles", _tono_signo),
            _col("variacion_pct", "Variación (%)", "pct", _tono_signo),
        ],
        "orden": ("variacion_valor", "desc"),
        "filtros": {"f_material": "material_codigo"},
        "resumen": _resumen_variacion,
    },
    "precios": {
        "titulo": "Comparaciones de precio (material y mes)",
        "nota": "Una fila por material, centro y mes con compra: precio real del mes (importe de las entradas 101/102 + "
                "ajustes de valor de la FBL3N, ÷ cantidad) contra el precio base. Bolsa = (precio − precio base) × cantidad. Las atípicas no suman al IPU.",
        "columnas": [
            _col("mes_texto", "Mes"), _COL_MATERIAL, _COL_DESCRIPCION, _col("centro", "Centro", "centro"),
            _col("umb", "UMB"), _col("toneladas", "Cantidad", "cantidad"), _col("importe", "Importe (S/)", "soles"),
            _col("precio", "Precio", "numero"), _col("base_texto", "Comparado contra"),
            _col("precio_base", "Precio base", "numero"),
            _col("variacion_pct", "Variación", "pct", lambda v: _tono_signo(v, inverso=True)),
            _col("bolsa", "Bolsa (S/)", "soles", lambda v: _tono_signo(v, inverso=True)),
            _col("atipica_texto", "Atípica"),
        ],
        "orden": ("bolsa", "desc"),
        "filtros": {"f_material": "material_codigo", "f_centro": "centro", "f_mes_ipu": "mes_texto"},
        "resumen": _resumen_precios,
    },
    "entradas": {
        "titulo": "Entradas de mercancía de compra (MB51)",
        "nota": "Una fila por entrada 101 (o su anulación 102) con proveedor, de los materiales y meses comparados: de "
                "aquí sale el precio MB51 de cada mes (importe ÷ cantidad); el precio real le suma los ajustes de la "
                "FBL3N (ver el detalle por material en Planificación > Precios de entrada).",
        "columnas": [
            _col("fecha", "Fecha", "fecha"), _col("documento", "Doc. material", "codigo"),
            _col("posicion", "Pos.", "codigo"), _col("clase_movimiento", "Cl. mov.", "codigo"),
            _COL_MATERIAL, _COL_DESCRIPCION, _col("centro", "Centro", "centro"),
            _col("proveedor", "Proveedor", "codigo"), _col("pedido", "Pedido", "codigo"),
            _col("cantidad", "Cantidad", "cantidad"), _col("unidad", "UM"),
            _col("importe", "Importe (S/)", "soles"), _col("precio", "Precio unitario", "numero"),
        ],
        "orden": ("fecha", "desc"),
        "filtros": {"f_material": "material_codigo", "f_centro": "centro"},
        "resumen": _resumen_entradas,
    },
    "abastecimiento": {
        "titulo": "Plan vs real por material",
        "nota": "Plan de Abastecimiento S&OP (última versión) contra los ingresos de báscula (ZMM056, Recepción) del mes, "
                "por material y centro.",
        "columnas": [
            _COL_MATERIAL, _COL_DESCRIPCION, _col("centro", "Centro", "centro"), _col("familia", "Familia"),
            _col("cantidad_plan", "Plan (T)", "entero"), _col("cantidad_real", "Real (T)", "entero"),
            _col("diferencia", "Diferencia (T)", "entero", _tono_signo), _col("cumplimiento_pct", "% Cumplimiento", "pct"),
        ],
        "orden": ("cantidad_plan", "desc"),
        "filtros": {"f_centro": "centro", "f_material": "material_codigo"},
        "resumen": _resumen_abastecimiento,
    },
    "ingresos": {
        "titulo": "Tickets de báscula (ZMM056)",
        "nota": "Una fila por ticket de Recepción del mes (Tipo 1: terceros; Tipo 5: propio, cantera o traslado). "
                "En los propios, el proveedor es la planta o cantera de origen.",
        "columnas": [
            _col("fecha", "Fecha", "fecha"), _col("ticket", "Ticket", "codigo"), _col("centro", "Centro", "centro"),
            _COL_MATERIAL, _COL_DESCRIPCION, _col("proveedor", "Proveedor u origen"), _col("tipo", "Tipo"),
            _col("toneladas", "Toneladas", "numero"), _col("documento", "Doc. material", "codigo"),
        ],
        "orden": ("fecha", "desc"),
        "filtros": {"f_centro": "centro", "f_material": "material_codigo", "f_proveedor": "proveedor"},
        "resumen": _resumen_ingresos,
    },
}

# Títulos de los filtros de período (las demás chips usan el título de su columna).
_TITULOS_FILTRO = {"f_mes": "Mes", "f_anio": "Año", "f_contrato": "Bajo contrato", "f_unico": "Un solo proveedor",
                   "f_clase_mov": "Clase de movimiento", "f_mes_ipu": "Mes"}


# --- Memoria de las últimas consultas ---------------------------------------
# Paginar, ordenar o buscar no repite la consulta: se guardan las filas base
# de los últimos detalles abiertos (por versión de datos).
_MAX_EN_MEMORIA = 12
_memoria = collections.OrderedDict()
_candado = threading.Lock()


def _base(clave_cache, calcular, definicion):
    clave = (q._version_datos(), clave_cache)
    with _candado:
        if clave in _memoria:
            _memoria.move_to_end(clave)
            return _memoria[clave]
    df, info = calcular()
    df = df.reset_index(drop=True)
    # Texto de búsqueda de cada fila: todas sus columnas de texto juntas.
    texto = [c["clave"] for c in definicion["columnas"] if c["tipo"] in _TIPOS_TEXTO and c["clave"] in df]
    buscar = pd.Series("", index=df.index)
    for columna in texto:
        buscar = buscar + " " + df[columna].fillna("").astype(str)
    df["_buscar"] = buscar.str.lower()
    with _candado:
        _memoria[clave] = (df, info)
        while len(_memoria) > _MAX_EN_MEMORIA:
            _memoria.popitem(last=False)
    return df, info


# --- Formato ----------------------------------------------------------------

def _vacio(valor):
    if valor is None:
        return True
    try:
        return bool(pd.isna(valor))
    except (TypeError, ValueError):
        return False


def _formato(valor, tipo):
    if _vacio(valor) or valor == "":
        return "—"
    if tipo == "fecha":
        return valor.strftime("%d/%m/%Y") if hasattr(valor, "strftime") else str(valor)[:10]
    if tipo == "soles":
        return f"S/ {float(valor):z,.2f}"
    if tipo == "numero":
        return f"{float(valor):z,.2f}"
    if tipo == "cantidad":
        texto = f"{float(valor):z,.3f}".rstrip("0").rstrip(".")
        return texto if texto not in ("", "-0") else "0"
    if tipo in ("entero", "dias"):
        return f"{float(valor):z,.0f}"
    if tipo == "pct":
        return f"{float(valor):z,.1f}%"
    if tipo == "centro":
        return q.nombre_centro(valor)
    return str(valor)


def _celdas(df, columnas):
    filas = []
    for registro in df[[c["clave"] for c in columnas]].itertuples(index=False):
        celdas = []
        for c, valor in zip(columnas, registro):
            clase = "num" if c["tipo"] in _TIPOS_NUMERICOS else ("col-codigo" if c["tipo"] == "codigo" else "")
            tono = c["tono"](valor) if c["tono"] and not _vacio(valor) else ""
            celdas.append({"texto": _formato(valor, c["tipo"]), "clase": f"{clase} {tono}".strip()})
        filas.append(celdas)
    return filas


def _valor_excel(valor, tipo):
    if _vacio(valor) or valor == "":
        return None
    if tipo == "centro":
        return q.nombre_centro(valor)
    if tipo in _TIPOS_NUMERICOS:
        return round(float(valor), 2) if tipo == "pct" else float(valor)
    if tipo == "fecha":
        return valor if isinstance(valor, datetime.date) else str(valor)[:10]
    return str(valor)


def _excel(df, columnas, nombre):
    """Excel con todas las filas filtradas. openpyxl en modo de solo
    escritura (fila a fila): con pandas.to_excel, 70 mil movimientos de
    consumo tardaban ~43 s."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    libro = Workbook(write_only=True)
    hoja = libro.create_sheet("Detalle")
    hoja.freeze_panes = "A2"
    for i, c in enumerate(columnas, start=1):
        ancho = 40 if c["clave"] in ("texto_breve", "proveedor", "nombres", "motivo", "texto_clase") else 14
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.append([c["titulo"] for c in columnas])
    tipos = [c["tipo"] for c in columnas]
    for registro in df[[c["clave"] for c in columnas]].itertuples(index=False):
        hoja.append([_valor_excel(v, t) for v, t in zip(registro, tipos)])
    buffer = io.BytesIO()
    libro.save(buffer)
    respuesta = HttpResponse(buffer.getvalue(),
                             content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return respuesta


# --- Respuesta --------------------------------------------------------------

def responder(request, clave, calcular, clave_cache, nombre_archivo):
    """HTML del detalle (cuerpo completo, o solo las filas siguientes con
    ?_parte=filas) o el Excel con todas las filas (?formato=xlsx)."""
    definicion = DESGLOSES.get(clave)
    if definicion is None:
        raise Http404("Detalle no encontrado")
    columnas = definicion["columnas"]
    base, info = _base(clave_cache, calcular, definicion)
    df = base

    # Filtros por clic (chips quitables): igualdad contra una columna.
    chips = list(info.get("chips", []))
    for parametro, columna in definicion["filtros"].items():
        valor = request.GET.get(parametro)
        if not valor or columna not in df:
            continue
        df = df[df[columna].astype(str) == valor]
        col = next((c for c in columnas if c["clave"] == columna), None)
        titulo = _TITULOS_FILTRO.get(parametro) or (col["titulo"] if col else columna)
        texto = _formato(valor, col["tipo"]) if col and col["tipo"] == "centro" else valor
        chips.append({"parametro": parametro, "texto": f"{titulo}: {texto}"})

    buscar = request.GET.get("q", "").strip()
    if buscar and "_buscar" in df and not df.empty:
        df = df[df["_buscar"].str.contains(buscar.lower(), regex=False, na=False)]

    claves = {c["clave"]: c for c in columnas}
    orden = request.GET.get("orden")
    if orden not in claves:
        orden = definicion["orden"][0]
        direccion = definicion["orden"][1]
    else:
        direccion = "asc" if request.GET.get("dir") == "asc" else "desc"
    if not df.empty and orden in df:
        clave_orden = None
        if claves[orden]["tipo"] in _TIPOS_TEXTO:
            clave_orden = lambda s: s.fillna("").astype(str).str.lower()  # noqa: E731
        elif claves[orden]["tipo"] == "fecha":
            clave_orden = lambda s: pd.to_datetime(s, errors="coerce")  # noqa: E731
        df = df.sort_values(orden, ascending=direccion == "asc", na_position="last", key=clave_orden, kind="stable")

    if request.GET.get("formato") == "xlsx":
        return _excel(df.head(MAX_FILAS_EXCEL), columnas, nombre_archivo)

    try:
        offset = max(0, int(request.GET.get("offset", 0)))
    except ValueError:
        offset = 0
    pagina = df.iloc[offset:offset + FILAS_POR_PAGINA]
    total = len(df)
    hay_mas = offset + FILAS_POR_PAGINA < total
    filas = _celdas(pagina, columnas)

    if request.GET.get("_parte") == "filas":
        respuesta = HttpResponse(render_to_string("gestion/_desglose_filas.html", {"filas": filas}))
        respuesta["X-Hay-Mas"] = "true" if hay_mas else "false"
        respuesta["X-Total"] = str(total)
        return respuesta

    parametros = request.GET.copy()
    for p in ("offset", "_parte", "formato"):
        parametros.pop(p, None)
    parametros["formato"] = "xlsx"
    contexto = {
        "definicion": definicion, "columnas": columnas, "filas": filas, "total": total, "base_total": len(base),
        "mostradas": min(offset + FILAS_POR_PAGINA, total), "hay_mas": hay_mas, "por_pagina": FILAS_POR_PAGINA,
        "orden": orden, "direccion": direccion, "buscar": buscar, "chips": chips,
        "resumen": definicion["resumen"](df) if not df.empty else [],
        "periodo": info.get("periodo", ""), "aviso": info.get("aviso", ""),
        "url_excel": f"{request.path}?{parametros.urlencode()}", "excel_tope": total > MAX_FILAS_EXCEL,
        "max_excel": MAX_FILAS_EXCEL,
    }
    return HttpResponse(render_to_string("gestion/_desglose_cuerpo.html", contexto, request=request))


def filtro(**valores):
    """Valor de data-filtrar (querystring de filtros por clic) para las plantillas/vistas."""
    return urlencode({k: v for k, v in valores.items() if v not in (None, "")})


def mes_param(fecha):
    """f_mes de un mes: '2026-03'."""
    return f"{fecha.year}-{fecha.month:02d}"


def leer_mes(texto):
    """Primer y último día del mes de un f_mes ('2026-03'), o None si no es válido."""
    try:
        anio, mes = (int(x) for x in str(texto).split("-")[:2])
        inicio = datetime.date(anio, mes, 1)
    except (TypeError, ValueError):
        return None
    if not 2000 <= anio <= 2100:
        return None
    siguiente = datetime.date(anio + (mes == 12), mes % 12 + 1, 1)
    return inicio, siguiente - datetime.timedelta(days=1)
